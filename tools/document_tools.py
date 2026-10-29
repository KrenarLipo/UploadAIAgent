"""Tool that renders the food guide as PDF and DOCX."""

import re
from datetime import date
from pathlib import Path
from typing import TypedDict
from xml.sax.saxutils import escape

from anthropic import beta_tool

OUTPUT_DIR = Path("output")


class Table(TypedDict):
    columns: list[str]
    rows: list[list[str]]


class Section(TypedDict, total=False):
    heading: str
    paragraphs: list[str]
    bullets: list[str]
    table: Table


def _slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-") or "food-guide"


def _write_pdf(path: Path, title: str, subtitle: str, sections: list[Section]) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table as RLTable, TableStyle,
    )

    styles = getSampleStyleSheet()
    cell = styles["BodyText"].clone("cell", fontSize=9, leading=11)
    story = [Paragraph(escape(title), styles["Title"])]
    if subtitle:
        story.append(Paragraph(escape(subtitle), styles["Italic"]))
    story.append(Spacer(1, 0.5 * cm))

    for s in sections:
        story.append(Paragraph(escape(s.get("heading", "")), styles["Heading2"]))
        for para in s.get("paragraphs", []):
            story.append(Paragraph(escape(para), styles["BodyText"]))
        if s.get("bullets"):
            story.append(ListFlowable(
                [ListItem(Paragraph(escape(b), styles["BodyText"])) for b in s["bullets"]],
                bulletType="bullet",
            ))
        if s.get("table"):
            t = s["table"]
            data = [[Paragraph(f"<b>{escape(c)}</b>", cell) for c in t["columns"]]]
            data += [[Paragraph(escape(str(v)), cell) for v in row] for row in t["rows"]]
            table = RLTable(data, repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2e7d32")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f8e9")]),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            story += [Spacer(1, 0.2 * cm), table]
        story.append(Spacer(1, 0.4 * cm))

    SimpleDocTemplate(str(path), pagesize=A4, title=title,
                      leftMargin=2 * cm, rightMargin=2 * cm).build(story)


def _write_docx(path: Path, title: str, subtitle: str, sections: list[Section]) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading(title, level=0)
    if subtitle:
        doc.add_paragraph(subtitle).italic = True
    for s in sections:
        doc.add_heading(s.get("heading", ""), level=1)
        for para in s.get("paragraphs", []):
            doc.add_paragraph(para)
        for b in s.get("bullets", []):
            doc.add_paragraph(b, style="List Bullet")
        if s.get("table"):
            t = s["table"]
            table = doc.add_table(rows=1, cols=len(t["columns"]))
            table.style = "Light Grid Accent 1"
            for c, name in zip(table.rows[0].cells, t["columns"]):
                c.text = name
            for row in t["rows"]:
                for c, v in zip(table.add_row().cells, row):
                    c.text = str(v)
    doc.save(str(path))


@beta_tool
def generate_food_guide(title: str, sections: list[Section], subtitle: str = "") -> str:
    """Render the food guide to both PDF and Word (.docx) in the output folder.

    Call this once the content is ready. Each section has a heading and any mix of
    paragraphs, bullets and a table (columns plus rows of string cells, one cell per column).

    Args:
        title: Guide title shown on the first page.
        sections: Ordered list of sections making up the guide.
        subtitle: Optional line under the title, e.g. audience or date.
    """
    for i, s in enumerate(sections):
        t = s.get("table")
        if t and any(len(r) != len(t["columns"]) for r in t["rows"]):
            return f"Error: section {i} ('{s.get('heading')}') has table rows whose length differs from its columns."

    OUTPUT_DIR.mkdir(exist_ok=True)
    stem = OUTPUT_DIR / f"{_slug(title)}-{date.today().isoformat()}"
    pdf, docx = stem.with_suffix(".pdf"), stem.with_suffix(".docx")
    _write_pdf(pdf, title, subtitle, sections)
    _write_docx(docx, title, subtitle, sections)
    return f"Created:\n{pdf.resolve()}\n{docx.resolve()}"
