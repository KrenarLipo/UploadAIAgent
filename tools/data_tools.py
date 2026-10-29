"""Tools for finding and reading the input files."""

import csv
import json
from pathlib import Path

from anthropic import beta_tool

SUPPORTED_DATA = {".csv", ".tsv", ".xlsx", ".xlsm", ".json"}


@beta_tool
def list_files(directory: str, extension: str = "") -> str:
    """List files in a directory, e.g. to find the PDFs to upload or the data file to use.

    Args:
        directory: Directory to list, relative to the project folder or absolute.
        extension: Optional filter such as ".pdf" or ".csv". Empty lists everything.
    """
    root = Path(directory).expanduser()
    if not root.is_dir():
        return f"Error: {root} is not a directory."
    files = sorted(
        p for p in root.rglob("*")
        if p.is_file() and (not extension or p.suffix.lower() == extension.lower())
    )
    if not files:
        return f"No files found in {root}."
    return "\n".join(f"{p}  ({p.stat().st_size:,} bytes)" for p in files)


@beta_tool
def read_data_file(path: str, sheet: str = "") -> str:
    """Read a structured data file (CSV, TSV, Excel or JSON) and return its rows as JSON.

    Args:
        path: Path to the data file.
        sheet: For Excel files, the sheet name to read. Empty reads the first sheet.
    """
    p = Path(path).expanduser()
    if not p.is_file():
        return f"Error: {p} does not exist."
    suffix = p.suffix.lower()
    if suffix not in SUPPORTED_DATA:
        return f"Error: unsupported file type {suffix}. Supported: {', '.join(sorted(SUPPORTED_DATA))}."

    if suffix in {".csv", ".tsv"}:
        with p.open(newline="", encoding="utf-8-sig") as f:
            sample = f.read(4096)
            f.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel_tab if suffix == ".tsv" else csv.excel
            rows = list(csv.DictReader(f, dialect=dialect))
    elif suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        wb = load_workbook(p, read_only=True, data_only=True)
        ws = wb[sheet] if sheet else wb.worksheets[0]
        values = list(ws.iter_rows(values_only=True))
        if not values:
            return json.dumps({"sheets": wb.sheetnames, "columns": [], "rows": []})
        header = [str(h) if h is not None else f"column_{i + 1}" for i, h in enumerate(values[0])]
        rows = [
            {k: v for k, v in zip(header, r)}
            for r in values[1:]
            if any(v is not None for v in r)
        ]
        return json.dumps(
            {"sheets": wb.sheetnames, "columns": header, "row_count": len(rows), "rows": rows},
            default=str, ensure_ascii=False,
        )
    else:
        data = json.loads(p.read_text(encoding="utf-8"))
        return json.dumps(data, ensure_ascii=False)

    columns = list(rows[0].keys()) if rows else []
    return json.dumps(
        {"columns": columns, "row_count": len(rows), "rows": rows}, ensure_ascii=False
    )
